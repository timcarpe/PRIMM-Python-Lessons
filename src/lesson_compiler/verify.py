"""Verify a compiled Python lesson suite."""

from __future__ import annotations

import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from lesson_compiler.paths import (
    CONFIG_PATH,
    DEFAULT_REPOSITORY_ROOT,
    LEARNER_DIRECTORY,
    RECORDS_ROOT,
    SUPPLEMENTAL_DIRECTORY,
)
from lesson_compiler.support import patch_standard_record, write_manifest


def plain(text: str) -> str:
    """Normalize prose for cross-format challenge-text comparisons."""
    text = re.sub(r"-\s*\n\s*", "-", text)
    return re.sub(r"\s+", " ", text).strip()


def strip_pdf_page_furniture(text: str) -> str:
    """Remove repeated worksheet headers and footers from extracted page text."""
    return "\n".join(
        line
        for line in text.splitlines()
        if not re.fullmatch(r"\s*\d+\s*", line)
        and not re.fullmatch(r"\s*© Tim Carpenter \d{4}\s*", line)
    )


def verify_suite(
    output_root: Path = DEFAULT_REPOSITORY_ROOT,
    work_root: Path | None = None,
) -> dict:
    """Verify inventory, content contracts, and key formatting contracts."""
    output_root = output_root.resolve()
    work_root = work_root.resolve() if work_root is not None else None
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    expected = set(config["lesson_numbers"])
    changes = {
        tuple(int(value) for value in key.split(".")): text
        for key, text in config["challenge_changes"].items()
    }
    failures: list[str] = []
    if len(changes) != config["expected_challenge_change_count"]:
        failures.append(
            "canonical challenge-change count mismatch: "
            f"expected {config['expected_challenge_change_count']}, "
            f"found {len(changes)}"
        )

    stats = {
        "lessons": 0,
        "pptx": 0,
        "pdf": 0,
        "docx": 0,
        "python": 0,
        "slide_pages": 0,
    }
    learner = output_root / LEARNER_DIRECTORY
    supplemental = output_root / SUPPLEMENTAL_DIRECTORY
    actual = {int(path.name.split()[1]) for path in learner.glob("Grade */Lesson *")}
    if actual != expected:
        failures.append(
            f"lesson set mismatch: expected {sorted(expected)}, got {sorted(actual)}"
        )

    for number in sorted(expected):
        record = patch_standard_record(
            json.loads(
                (RECORDS_ROOT / f"lesson{number:02d}.json").read_text(
                    encoding="utf-8"
                )
            ),
            changes,
            config,
        )
        learner_dir = next(learner.glob(f"Grade */Lesson {number:02d} - *"), None)
        supplemental_dir = next(
            supplemental.glob(f"Grade */Lesson {number:02d} - *"),
            None,
        )
        if learner_dir is None or supplemental_dir is None:
            failures.append(f"L{number:02d}: missing learner or supplemental folder")
            continue
        stats["lessons"] += 1
        presentations = list(learner_dir.glob("* - Slides.pptx"))
        challenge_pdfs = list(learner_dir.glob("* - Challenges.pdf"))
        starters = list(learner_dir.glob("* - Starter Code.py"))
        documents = list(supplemental_dir.glob("*.docx"))
        programs = list((supplemental_dir / "Solutions").glob("*.py"))
        inventory = [
            len(presentations),
            len(challenge_pdfs),
            len(starters),
            len(documents),
            len(programs),
        ]
        if inventory != [1, 1, 1, 2, 5]:
            failures.append(f"L{number:02d}: inventory {inventory}")
            continue
        stats["pptx"] += 1
        stats["pdf"] += 1
        stats["docx"] += 2
        stats["python"] += 6

        presentation = Presentation(presentations[0])
        slide_count = len(presentation.slides)
        stats["slide_pages"] += slide_count
        if slide_count != 9:
            failures.append(f"L{number:02d}: {slide_count} slides, expected 9")
        deck_text = "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        lowered = deck_text.lower()
        if "suggested solution" in lowered:
            failures.append(f"L{number:02d}: solution slide text in learner deck")
        challenge_reader = PdfReader(str(challenge_pdfs[0]))
        challenge_page_count = len(challenge_reader.pages)
        if challenge_page_count not in {1, 2}:
            failures.append(
                f"L{number:02d}: challenge PDF has {challenge_page_count} pages; "
                "expected one or two"
            )

        worksheet = next(supplemental_dir.glob("* - Worksheet.docx"))
        worksheet_text = "\n".join(
            paragraph.text for paragraph in Document(worksheet).paragraphs
        )
        pdf_text = "\n".join(
            strip_pdf_page_furniture(page.extract_text() or "")
            for page in challenge_reader.pages
        )
        prompts = record["worksheet"]["page2"]["challenges"]
        for challenge, item in enumerate(prompts, 1):
            replacement = item["prompt"]
            needle = plain(replacement)
            for artifact, text in (
                ("deck", deck_text),
                ("worksheet", worksheet_text),
                ("PDF", pdf_text),
            ):
                if needle not in plain(text):
                    failures.append(
                        f"L{number:02d} C{challenge}: accepted prompt absent "
                        f"from {artifact}"
                    )

        if work_root is not None:
            rendered = (
                work_root / f"renders/documents/lesson{number:02d}/{worksheet.stem}.pdf"
            )
            if not rendered.exists():
                failures.append(
                    f"L{number:02d}: rendered worksheet source is missing"
                )
            else:
                worksheet_page_count = len(PdfReader(str(rendered)).pages)
                if worksheet_page_count not in {3, 4}:
                    failures.append(
                        f"L{number:02d}: rendered worksheet source has "
                        f"{worksheet_page_count} pages; expected three or four"
                    )
                expected_challenge_pages = worksheet_page_count - 2
                if challenge_page_count != expected_challenge_pages:
                    failures.append(
                        f"L{number:02d}: challenge PDF has "
                        f"{challenge_page_count} pages; expected "
                        f"{expected_challenge_pages} from worksheet"
                    )

    lesson_one = next(supplemental.glob("Grade 6/Lesson 01 -*/* - Worksheet.docx"))
    with ZipFile(lesson_one) as archive:
        lesson_one_xml = archive.read("word/document.xml").decode("utf-8")
    if "<wp:anchor" not in lesson_one_xml or "<wp:inline" in lesson_one_xml:
        failures.append("L01: worksheet illustration is not anchored")
    if not re.search(r"<wp:wrap(?:Square|Tight|Through)\b", lesson_one_xml):
        failures.append("L01: worksheet illustration has no text-wrapping rule")

    report = {
        "passed": not failures,
        "stats": stats,
        "excluded_capstones": config["excluded_capstones"],
        "failures": failures,
    }
    (output_root / "VERIFICATION.json").write_text(
        json.dumps(report, indent=2) + "\n",
        encoding="utf-8",
    )
    write_manifest(output_root, config)
    return report


def main() -> None:
    """Verify the canonical root lesson folders."""
    report = verify_suite()
    print(json.dumps(report, indent=2))
    raise SystemExit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
