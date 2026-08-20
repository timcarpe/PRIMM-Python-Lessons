from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = ROOT / "lesson_compiler/curriculum/recompilation_run_2026-08-20.json"
DEFAULT_OUTPUT = ROOT / "build/Recompiled Lesson Suite - Non-Capstone"
DEFAULT_WORK = ROOT / "build/work"
EXPECTED = {number for number in range(1, 37) if number not in {11, 27, 36}}


def plain(text: str) -> str:
    return re.sub(
        r"\s+",
        " ",
        text.replace('"', "").replace("“", "").replace("”", "").replace("'", ""),
    ).strip()


def main() -> None:
    parser = argparse.ArgumentParser(description="Verify the accepted non-capstone lesson build")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--work", type=Path, default=DEFAULT_WORK)
    args = parser.parse_args()
    output = args.output.resolve()
    work = args.work.resolve()

    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    changes = {
        tuple(int(value) for value in key.split(".")): text
        for key, text in config["challenge_changes"].items()
    }
    failures: list[str] = []
    stats = {
        "lessons": 0,
        "pptx": 0,
        "pdf": 0,
        "docx": 0,
        "python": 0,
        "slide_pages": 0,
    }
    learner = output / "Mapped Curriculum Lessons"
    supplemental = output / "Mapped Curriculum Lesson - Supplemental Material"
    actual = {int(path.name.split()[1]) for path in learner.glob("Grade */Lesson *")}
    if actual != EXPECTED:
        failures.append(f"lesson set mismatch: expected {sorted(EXPECTED)}, got {sorted(actual)}")

    for number in sorted(EXPECTED):
        learner_dir = next(learner.glob(f"Grade */Lesson {number:02d} - *"), None)
        supp_dir = next(supplemental.glob(f"Grade */Lesson {number:02d} - *"), None)
        if learner_dir is None or supp_dir is None:
            failures.append(f"L{number:02d}: missing learner or supplemental folder")
            continue
        stats["lessons"] += 1
        pptx = list(learner_dir.glob("* - Slides.pptx"))
        pdf = list(learner_dir.glob("* - Challenges.pdf"))
        starter = list(learner_dir.glob("* - Starter Code.py"))
        docx = list(supp_dir.glob("*.docx"))
        py = list((supp_dir / "Solutions").glob("*.py"))
        inventory = [len(pptx), len(pdf), len(starter), len(docx), len(py)]
        if inventory != [1, 1, 1, 2, 5]:
            failures.append(f"L{number:02d}: inventory {inventory}")
            continue
        stats["pptx"] += 1
        stats["pdf"] += 1
        stats["docx"] += 2
        stats["python"] += 6

        presentation = Presentation(pptx[0])
        slide_count = len(presentation.slides)
        stats["slide_pages"] += slide_count
        expected_slides = 13 if number == 1 else 9
        if slide_count != expected_slides:
            failures.append(f"L{number:02d}: {slide_count} slides, expected {expected_slides}")
        deck_text = "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        if number != 1:
            lowered = deck_text.lower()
            if "suggested solution" in lowered or "fresh-copy" in lowered or "fresh start" in lowered:
                failures.append(f"L{number:02d}: forbidden solution/fresh text in learner deck")
        if len(PdfReader(str(pdf[0])).pages) != 1:
            failures.append(f"L{number:02d}: challenge PDF is not one page")

        worksheet = next(supp_dir.glob("* - Worksheet.docx"))
        worksheet_text = "\n".join(paragraph.text for paragraph in Document(worksheet).paragraphs)
        pdf_text = PdfReader(str(pdf[0])).pages[0].extract_text() or ""
        for challenge in range(1, 4):
            replacement = changes.get((number, challenge))
            if not replacement:
                continue
            needle = plain(replacement)
            for artifact, text in (
                ("deck", deck_text),
                ("worksheet", worksheet_text),
                ("PDF", pdf_text),
            ):
                if needle not in plain(text):
                    failures.append(
                        f"L{number:02d} C{challenge}: accepted prompt absent from {artifact}"
                    )

        if number != 1:
            render_pdf = work / f"renders/documents/lesson{number:02d}/{worksheet.stem}.pdf"
            if not render_pdf.exists() or len(PdfReader(str(render_pdf)).pages) != 3:
                failures.append(f"L{number:02d}: rendered worksheet source is not three pages")

    l01_source = ROOT / "Example Output/Lesson Suite/Lesson 01 - Hello Python!"
    l01_learner = learner / "Grade 6/Lesson 01 - Hello Python!"
    l01_supp = supplemental / "Grade 6/Lesson 01 - Hello Python!"
    for source_path in sorted(path for path in l01_source.rglob("*") if path.is_file()):
        relative = source_path.relative_to(l01_source)
        target = l01_supp / relative if relative.parts[0] == "Solutions" or source_path.suffix == ".docx" else l01_learner / relative
        if not target.exists() or source_path.read_bytes() != target.read_bytes():
            failures.append(f"L01: preserved original mismatch: {relative.as_posix()}")

    l02_pptx = next(learner.glob("Grade 6/Lesson 02 -*/* - Slides.pptx"))
    with ZipFile(l02_pptx) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide[678]\.xml", name)
        )
    if "&quot;Total:&quot;" not in xml and '"Total:"' not in xml:
        failures.append('L02: quoted "Total:" absent from challenge slides')
    if "067D17" not in xml:
        failures.append("L02: green literal colour absent from challenge slides")
    l02_docx = next(supplemental.glob("Grade 6/Lesson 02 -*/* - Worksheet.docx"))
    with ZipFile(l02_docx) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    if "067D17" not in xml or "Total:" not in xml:
        failures.append("L02: quoted green literal absent from worksheet")

    report = {
        "passed": not failures,
        "stats": stats,
        "preserved_original_lessons": [1],
        "excluded_capstones": [11, 27, 36],
        "failures": failures,
    }
    report_path = output / "VERIFICATION.json"
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))
    raise SystemExit(bool(failures))


if __name__ == "__main__":
    main()
